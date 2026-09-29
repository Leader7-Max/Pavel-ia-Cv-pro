"""Pavel CV ia - Étape 1 : socle (accueil, création de CV, lettre, export PDF/Word)."""
import io
import json

import streamlit as st
from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from google import genai

st.set_page_config(page_title="Pavel CV ia", page_icon="📄", layout="centered")

MODEL = "gemini-2.5-flash"
PAYS = ["France", "Suisse", "Belgique", "Canada", "Luxembourg", "Allemagne", "Royaume-Uni"]
LANGUES = ["Français", "English", "Deutsch", "Español", "Italiano"]
MODELES = {"Classic": (30, 30, 30), "Modern": (25, 80, 160), "ATS": (0, 0, 0)}


# ---------------------------------------------------------------- utilitaires
def clean(text):
    """fpdf (polices de base) ne gère que le latin-1 : on remplace le reste."""
    table = {"’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "—": "-",
             "•": "-", "…": "...", "€": "EUR", "\u00a0": " "}
    text = str(text or "")
    for old, new in table.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


def ask_gemini(prompt):
    try:
        key = st.secrets["GEMINI_API_KEY"]
    except Exception:
        st.error("Clé manquante : ajoute GEMINI_API_KEY dans les Secrets Streamlit.")
        st.stop()
    try:
        client = genai.Client(api_key=key)
        response = client.models.generate_content(model=MODEL, contents=prompt)
        return response.text or ""
    except Exception as exc:
        st.error(f"Erreur Gemini : {exc}")
        return ""


def parse_json(text):
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        return None
    try:
        return json.loads(text[start:end + 1])
    except json.JSONDecodeError:
        return None


def go(page):
    st.session_state.page = page
    st.rerun()


def init_state():
    st.session_state.setdefault("page", "home")
    st.session_state.setdefault("step", 1)
    st.session_state.setdefault("data", {})
    st.session_state.setdefault("cv", None)
    st.session_state.setdefault("letter", "")


# ------------------------------------------------------------------- exports
def contact_line(d):
    parts = [d.get("email"), d.get("tel"), d.get("ville")]
    return "  |  ".join(p for p in parts if p)


def cv_pdf(cv, d, template):
    r, g, b = MODELES[template]
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    name = f"{d.get('prenom', '')} {d.get('nom', '')}".strip()

    if template == "Modern":
        pdf.set_fill_color(r, g, b)
        pdf.rect(0, 0, 210, 40, "F")
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 22)
        pdf.set_xy(15, 8)
        pdf.cell(0, 10, clean(name))
        pdf.set_font("Helvetica", "", 12)
        pdf.set_xy(15, 20)
        pdf.cell(0, 6, clean(cv.get("titre", "")))
        pdf.set_font("Helvetica", "", 9)
        pdf.set_xy(15, 29)
        pdf.cell(0, 5, clean(contact_line(d)))
        pdf.set_y(46)
    else:
        pdf.set_text_color(r, g, b)
        pdf.set_font("Helvetica", "B", 20)
        pdf.cell(0, 9, clean(name), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 12)
        pdf.cell(0, 6, clean(cv.get("titre", "")), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(0, 5, clean(contact_line(d)), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    labels = cv.get("labels", {})

    def section(key, default):
        pdf.ln(3)
        pdf.set_text_color(r, g, b)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 7, clean(labels.get(key, default)).upper(), new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(r, g, b)
        pdf.line(15, pdf.get_y(), 195, pdf.get_y())
        pdf.ln(2)
        pdf.set_text_color(30, 30, 30)
        pdf.set_font("Helvetica", "", 10)

    def body(text, bold=False):
        pdf.set_font("Helvetica", "B" if bold else "", 10)
        pdf.multi_cell(0, 5.5, clean(text), new_x="LMARGIN", new_y="NEXT")

    if cv.get("resume"):
        section("profil", "Profil")
        body(cv["resume"])
    if cv.get("competences"):
        section("competences", "Compétences")
        body(", ".join(cv["competences"]))
    if cv.get("experiences"):
        section("experience", "Expérience professionnelle")
        for e in cv["experiences"]:
            head = " - ".join(x for x in [e.get("poste"), e.get("entreprise"), e.get("periode")] if x)
            body(head, bold=True)
            for m in e.get("missions", []):
                body("- " + m)
            pdf.ln(1.5)
    if cv.get("formations"):
        section("formation", "Formation")
        for f in cv["formations"]:
            body(" - ".join(x for x in [f.get("diplome"), f.get("etablissement"), f.get("periode")] if x))
    if cv.get("langues"):
        section("langues", "Langues")
        body(", ".join(cv["langues"]))
    if cv.get("centres_interet"):
        section("interets", "Centres d'intérêt")
        body(", ".join(cv["centres_interet"]))
    return bytes(pdf.output())


def cv_docx(cv, d):
    doc = Document()
    name = f"{d.get('prenom', '')} {d.get('nom', '')}".strip()
    doc.add_heading(name, 0)
    doc.add_paragraph(cv.get("titre", ""))
    doc.add_paragraph(contact_line(d))
    labels = cv.get("labels", {})

    def block(key, default, items):
        if items:
            doc.add_heading(labels.get(key, default), 1)
            for it in items:
                doc.add_paragraph(it)

    block("profil", "Profil", [cv.get("resume", "")] if cv.get("resume") else [])
    block("competences", "Compétences", [", ".join(cv.get("competences", []))] if cv.get("competences") else [])
    if cv.get("experiences"):
        doc.add_heading(labels.get("experience", "Expérience professionnelle"), 1)
        for e in cv["experiences"]:
            p = doc.add_paragraph()
            run = p.add_run(" - ".join(x for x in [e.get("poste"), e.get("entreprise"), e.get("periode")] if x))
            run.bold = True
            for m in e.get("missions", []):
                doc.add_paragraph(m, style="List Bullet")
    if cv.get("formations"):
        doc.add_heading(labels.get("formation", "Formation"), 1)
        for f in cv["formations"]:
            doc.add_paragraph(" - ".join(x for x in [f.get("diplome"), f.get("etablissement"), f.get("periode")] if x))
    block("langues", "Langues", [", ".join(cv.get("langues", []))] if cv.get("langues") else [])
    block("interets", "Centres d'intérêt", [", ".join(cv.get("centres_interet", []))] if cv.get("centres_interet") else [])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def letter_pdf(text):
    pdf = FPDF()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, 20)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, clean(text), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def letter_docx(text):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for para in text.split("\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- pages
def page_home():
    st.title("📄 Pavel CV ia")
    st.caption("Créez. Améliorez. Adaptez. Postulez.")
    st.write("Que souhaitez-vous faire ?")
    if st.button("📄 Créer mon CV", use_container_width=True, type="primary"):
        st.session_state.step = 1
        go("cv")
    if st.button("✉️ Ma lettre de motivation", use_container_width=True):
        go("letter")
    st.info("🔒 Vos informations servent uniquement à générer vos documents.")


def back_home():
    if st.button("← Accueil"):
        go("home")


def page_cv():
    back_home()
    st.title("📄 Créer mon CV")
    step = st.session_state.step
    data = st.session_state.data
    if step <= 3:
        st.progress(step / 3, text=f"Étape {step} / 3")

    if step == 1:
        with st.form("f1"):
            st.subheader("👤 Informations personnelles")
            prenom = st.text_input("Prénom", data.get("prenom", ""))
            nom = st.text_input("Nom", data.get("nom", ""))
            email = st.text_input("Email", data.get("email", ""))
            tel = st.text_input("Téléphone", data.get("tel", ""))
            ville = st.text_input("Ville", data.get("ville", ""))
            pays = st.selectbox("Pays visé", PAYS, index=PAYS.index(data.get("pays", "France")))
            poste = st.text_input("Poste recherché", data.get("poste", ""))
            if st.form_submit_button("Continuer →", use_container_width=True):
                if not prenom or not nom or not poste:
                    st.warning("Prénom, nom et poste recherché sont obligatoires.")
                else:
                    data.update(prenom=prenom, nom=nom, email=email, tel=tel,
                                ville=ville, pays=pays, poste=poste)
                    st.session_state.step = 2
                    st.rerun()

    elif step == 2:
        with st.form("f2"):
            st.subheader("💼 Votre parcours")
            st.caption("Écrivez simplement, l'IA reformule. Elle n'invente rien.")
            exp = st.text_area("Expériences (poste, entreprise, durée, ce que vous faisiez)",
                               data.get("exp", ""), height=160)
            formation = st.text_area("Formations / diplômes", data.get("formation", ""), height=100)
            comp = st.text_area("Compétences", data.get("comp", ""), height=80)
            langues = st.text_input("Langues parlées", data.get("langues", ""))
            col1, col2 = st.columns(2)
            back = col1.form_submit_button("← Retour", use_container_width=True)
            nxt = col2.form_submit_button("Continuer →", use_container_width=True)
            if back or nxt:
                data.update(exp=exp, formation=formation, comp=comp, langues=langues)
                st.session_state.step = 1 if back else 3
                st.rerun()

    elif step == 3:
        with st.form("f3"):
            st.subheader("🎨 Options")
            langue = st.selectbox("Langue du CV", LANGUES,
                                  index=LANGUES.index(data.get("langue_cv", "Français")))
            modele = st.selectbox("Modèle", list(MODELES), index=list(MODELES).index(data.get("modele", "Classic")))
            col1, col2 = st.columns(2)
            back = col1.form_submit_button("← Retour", use_container_width=True)
            gen = col2.form_submit_button("✨ Générer", use_container_width=True, type="primary")
            if back:
                st.session_state.step = 2
                st.rerun()
            if gen:
                data.update(langue_cv=langue, modele=modele)
                generate_cv(data)

    elif step == 4:
        show_cv_result()


def generate_cv(d):
    prompt = f"""Tu es un expert en recrutement. Rédige un CV en {d['langue_cv']} adapté aux usages de ce pays : {d['pays']}.
Poste visé : {d['poste']}
Expériences : {d.get('exp', '')}
Formations : {d.get('formation', '')}
Compétences : {d.get('comp', '')}
Langues : {d.get('langues', '')}

RÈGLES : reformule et professionnalise, mais N'INVENTE AUCUNE expérience, entreprise, diplôme, date ni chiffre absent des informations fournies.
Réponds UNIQUEMENT avec un objet JSON valide, sans texte autour, de la forme :
{{"titre": "", "resume": "", "competences": [""],
"experiences": [{{"poste": "", "entreprise": "", "periode": "", "missions": [""]}}],
"formations": [{{"diplome": "", "etablissement": "", "periode": ""}}],
"langues": [""], "centres_interet": [""],
"labels": {{"profil": "", "competences": "", "experience": "", "formation": "", "langues": "", "interets": ""}}}}
Les "labels" sont les titres de sections traduits en {d['langue_cv']}."""
    with st.spinner("Rédaction en cours..."):
        cv = parse_json(ask_gemini(prompt))
    if cv is None:
        st.error("La réponse de l'IA n'a pas pu être lue. Réessayez.")
        return
    st.session_state.cv = cv
    st.session_state.step = 4
    st.rerun()


def show_cv_result():
    cv, d = st.session_state.cv, st.session_state.data
    st.success("✅ Votre CV est prêt")
    st.subheader(f"{d.get('prenom', '')} {d.get('nom', '')}")
    st.caption(cv.get("titre", ""))
    cv["resume"] = st.text_area("Profil (modifiable)", cv.get("resume", ""), height=120)
    if cv.get("competences"):
        st.markdown("**Compétences :** " + ", ".join(cv["competences"]))
    for e in cv.get("experiences", []):
        st.markdown(f"**{e.get('poste', '')}** - {e.get('entreprise', '')} ({e.get('periode', '')})")
        for m in e.get("missions", []):
            st.markdown(f"- {m}")
    for f in cv.get("formations", []):
        st.markdown(f"🎓 {f.get('diplome', '')} - {f.get('etablissement', '')} ({f.get('periode', '')})")

    nom_fichier = f"CV_{d.get('prenom', '')}_{d.get('nom', '')}".replace(" ", "_")
    try:
        st.download_button("⬇️ Télécharger en PDF", cv_pdf(cv, d, d.get("modele", "Classic")),
                           f"{nom_fichier}.pdf", "application/pdf", use_container_width=True)
        st.download_button("⬇️ Télécharger en Word", cv_docx(cv, d), f"{nom_fichier}.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           use_container_width=True)
    except Exception as exc:
        st.error(f"Erreur d'export : {exc}")
    if st.button("✏️ Modifier mes informations", use_container_width=True):
        st.session_state.step = 1
        st.rerun()


def page_letter():
    back_home()
    st.title("✉️ Lettre de motivation")
    d = st.session_state.data
    with st.form("lettre"):
        prenom = st.text_input("Prénom", d.get("prenom", ""))
        nom = st.text_input("Nom", d.get("nom", ""))
        poste = st.text_input("Poste recherché", d.get("poste", ""))
        entreprise = st.text_input("Entreprise")
        ville = st.text_input("Ville", d.get("ville", ""))
        annonce = st.text_area("Annonce (collez-la ici, facultatif)", height=150)
        parcours = st.text_area("Votre parcours en quelques lignes", d.get("exp", ""), height=120)
        langue = st.selectbox("Langue", LANGUES)
        go_btn = st.form_submit_button("✨ Générer la lettre", use_container_width=True, type="primary")
    if go_btn:
        if not poste or not entreprise:
            st.warning("Poste et entreprise sont obligatoires.")
        else:
            d.update(prenom=prenom, nom=nom, poste=poste, ville=ville)
            prompt = f"""Rédige une lettre de motivation en {langue}, professionnelle et personnalisée.
Candidat : {prenom} {nom}, ville : {ville}
Poste : {poste} - Entreprise : {entreprise}
Annonce : {annonce or 'non fournie'}
Parcours : {parcours}
RÈGLES : n'invente aucune expérience ni diplôme absent du parcours. Adapte la lettre aux compétences demandées dans l'annonce.
Réponds uniquement avec le texte de la lettre, prêt à envoyer."""
            with st.spinner("Rédaction en cours..."):
                st.session_state.letter = ask_gemini(prompt)
    if st.session_state.letter:
        st.session_state.letter = st.text_area("Votre lettre (modifiable)", st.session_state.letter, height=350)
        try:
            st.download_button("⬇️ PDF", letter_pdf(st.session_state.letter), "lettre.pdf",
                               "application/pdf", use_container_width=True)
            st.download_button("⬇️ Word", letter_docx(st.session_state.letter), "lettre.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)
        except Exception as exc:
            st.error(f"Erreur d'export : {exc}")


# ---------------------------------------------------------------------- main
init_state()
{"home": page_home, "cv": page_cv, "letter": page_letter}[st.session_state.page]()
 votre document")
            if missing:
                st.caption(
                    "Mots-clés absents (à ajouter seulement s'ils correspondent à votre parcours) : "
                    + ", ".join(missing[:10])
                )
