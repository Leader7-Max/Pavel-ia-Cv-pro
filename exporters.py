"""Exports PDF et Word du document généré."""

import re
from io import BytesIO

import streamlit as st
from fpdf import FPDF

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    DOCX_AVAILABLE = True
except ImportError:  # python-docx absent : l'export Word est simplement désactivé
    DOCX_AVAILABLE = False


# ─────────────────────────────────────────────────────────────
# Export PDF
# ─────────────────────────────────────────────────────────────
INK = (18, 39, 43)
BODY = (38, 52, 56)
TEAL = (14, 107, 107)
AMBER = (242, 169, 59)
LINE = (220, 227, 229)

PDF_REPLACEMENTS = {
    "\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-",
    "\u2026": "...", "\u2022": "-", "\u00a0": " ", "\u202f": " ", "\u0153": "oe", "\u0152": "Oe",
    "\u20ac": "EUR", "\u2192": "->",
}


def pdf_safe(text: str) -> str:
    for src, dst in PDF_REPLACEMENTS.items():
        text = text.replace(src, dst)
    text = text.encode("latin-1", "replace").decode("latin-1")
    # Coupe les « mots » très longs (URL…) qui feraient planter multi_cell.
    return re.sub(r"(\S{60})(?=\S)", r"\1 ", text)


def is_heading(line: str) -> bool:
    core = line.rstrip(":").strip()
    letters = [c for c in core if c.isalpha()]
    return bool(letters) and core == core.upper() and len(core) <= 48 and not core.startswith("-")


@st.cache_data(show_spinner=False)
def build_pdf(text: str, is_cv: bool) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    right = pdf.w - pdf.r_margin
    header_step = 0

    for raw in text.split("\n"):
        line = pdf_safe(raw).strip()
        if not line:
            pdf.ln(3)
            continue

        if is_cv and header_step == 0:
            pdf.set_font("Helvetica", "B", 22)
            pdf.set_text_color(*INK)
            pdf.multi_cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
            header_step = 1
            continue

        if is_cv and header_step == 1:
            header_step = 2
            if not is_heading(line):
                pdf.set_font("Helvetica", "", 12)
                pdf.set_text_color(*TEAL)
                pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
                y = pdf.get_y() + 2
                pdf.set_draw_color(*AMBER)
                pdf.set_line_width(0.8)
                pdf.line(pdf.l_margin, y, pdf.l_margin + 40, y)
                pdf.ln(5)
                continue

        if is_cv and is_heading(line):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 11)
            pdf.set_text_color(*TEAL)
            pdf.multi_cell(0, 6, line.rstrip(":"), new_x="LMARGIN", new_y="NEXT")
            y = pdf.get_y() + 0.5
            pdf.set_draw_color(*LINE)
            pdf.set_line_width(0.3)
            pdf.line(pdf.l_margin, y, right, y)
            pdf.ln(2.5)
        elif line.startswith("- "):
            pdf.set_font("Helvetica", "", 10.5)
            pdf.set_text_color(*BODY)
            pdf.set_x(pdf.l_margin + 2)
            pdf.cell(5, 5.6, "-")
            pdf.multi_cell(0, 5.6, line[2:].strip(), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 10.5)
            pdf.set_text_color(*BODY)
            pdf.multi_cell(0, 5.6, line, new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


# ─────────────────────────────────────────────────────────────
# Export Word (.docx)
# ─────────────────────────────────────────────────────────────
BULLET_RE = re.compile(r"^[•\-\*]\s+")


def _rgb(color: tuple):
    return RGBColor(*color)


def _bottom_border(paragraph, color: str = "DCE3E5", size: int = 6) -> None:
    """Ajoute un filet sous un paragraphe (titres de section)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


@st.cache_data(show_spinner=False)
def build_docx(text: str, is_cv: bool) -> bytes:
    """Génère un document Word A4 mis en forme, modifiable dans Word / LibreOffice / Google Docs."""
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(BODY)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.1

    header_step = 0

    for raw in text.split("\n"):
        line = raw.strip()

        if not line:
            if not is_cv:  # lettre : les lignes vides séparent les paragraphes
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue

        if is_cv and header_step == 0:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = _rgb(INK)
            p.paragraph_format.space_after = Pt(0)
            header_step = 1
            continue

        if is_cv and header_step == 1:
            header_step = 2
            if not is_heading(line):
                p = doc.add_paragraph()
                run = p.add_run(line)
                run.font.size = Pt(12.5)
                run.font.color.rgb = _rgb(TEAL)
                p.paragraph_format.space_after = Pt(6)
                _bottom_border(p, color="F2A93B", size=12)
                continue

        if is_cv and is_heading(line):
            p = doc.add_paragraph()
            run = p.add_run(line.rstrip(":"))
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = _rgb(TEAL)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            _bottom_border(p)
        elif BULLET_RE.match(line):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(BULLET_RE.sub("", line))
            p.paragraph_format.space_after = Pt(1.5)
        else:
            p = doc.add_paragraph(line)
            if not is_cv:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(4)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
              
